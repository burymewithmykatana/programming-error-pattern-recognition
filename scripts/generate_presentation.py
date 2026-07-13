"""Generate the editable project PowerPoint from measured experiment outputs."""

from __future__ import annotations

import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "project_presentation.pptx"
FIGURES = ROOT / "reports" / "figures"

INK = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(13, 148, 136)
AMBER = RGBColor(217, 119, 6)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)


def read_metrics(name: str) -> dict[str, float]:
    path = ROOT / "reports" / "experiments" / name / "metrics.json"
    return json.loads(path.read_text(encoding="utf-8"))


def create_figures() -> tuple[Path, Path, Path, Path]:
    FIGURES.mkdir(parents=True, exist_ok=True)
    synthetic_svm = read_metrics("synthetic_svm")
    synthetic_codebert = read_metrics("synthetic_codebert")
    real_svm = read_metrics("codecontests_svm")
    real_codebert = read_metrics("codecontests_codebert_reduced")

    metrics_path = FIGURES / "all_model_metrics.png"
    labels = [
        "Synthetic\nSVM",
        "Synthetic\nCodeBERT",
        "CodeContests\nSVM",
        "CodeContests\nCodeBERT",
    ]
    accuracy = [
        synthetic_svm["accuracy"],
        synthetic_codebert["accuracy"],
        real_svm["accuracy"],
        real_codebert["accuracy"],
    ]
    macro_f1 = [
        synthetic_svm["macro_f1"],
        synthetic_codebert["macro_f1"],
        real_svm["macro_f1"],
        real_codebert["macro_f1"],
    ]
    x = range(len(labels))
    plt.figure(figsize=(9, 4.8))
    plt.bar([i - 0.18 for i in x], accuracy, width=0.36, label="Accuracy", color="#2563eb")
    plt.bar([i + 0.18 for i in x], macro_f1, width=0.36, label="Macro F1", color="#0d9488")
    plt.xticks(list(x), labels)
    plt.ylim(0, 1)
    plt.ylabel("Score")
    plt.legend()
    plt.tight_layout()
    plt.savefig(metrics_path, dpi=180)
    plt.close()

    confusion_paths = []
    for name, title in (
        ("synthetic_svm", "Synthetic SVM confusion matrix"),
        ("codecontests_svm", "CodeContests SVM confusion matrix"),
        ("codecontests_codebert_reduced", "CodeContests CodeBERT confusion matrix"),
    ):
        matrix = pd.read_csv(
            ROOT / "reports" / "experiments" / name / "confusion_matrix.csv",
            index_col=0,
        )
        path = FIGURES / f"{name}_confusion.png"
        plt.figure(figsize=(7.2, 5.4))
        plt.imshow(matrix.values, cmap="Blues")
        plt.colorbar()
        plt.xticks(range(len(matrix.columns)), matrix.columns, rotation=50, ha="right", fontsize=7)
        plt.yticks(range(len(matrix.index)), matrix.index, fontsize=7)
        plt.title(title)
        for row in range(len(matrix.index)):
            for column in range(len(matrix.columns)):
                plt.text(column, row, str(matrix.iat[row, column]), ha="center", va="center", fontsize=7)
        plt.tight_layout()
        plt.savefig(path, dpi=180)
        plt.close()
        confusion_paths.append(path)
    return metrics_path, confusion_paths[0], confusion_paths[1], confusion_paths[2]


def set_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = LIGHT


def add_text(slide, text: str, left: float, top: float, width: float, height: float, size: int = 22,
             color=INK, bold: bool = False, align=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return box


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_text(slide, title, 0.55, 0.35, 12.2, 0.55, size=27, color=INK, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 0.92, 12.0, 0.36, size=13, color=TEAL)


def add_bullets(slide, items: list[str], left: float = 0.8, top: float = 1.55,
                width: float = 11.8, height: float = 5.3, size: int = 20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, text in enumerate(items):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = text
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = INK
        p.space_after = Pt(12)
    return box


def add_slide(prs: Presentation, title: str, bullets: list[str], subtitle: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, title, subtitle)
    add_bullets(slide, bullets)
    return slide


def add_table(slide, rows: list[list[str]], left: float, top: float, width: float, height: float,
              header_color=BLUE) -> None:
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)).table
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = header_color if row_index == 0 else WHITE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(11 if len(value) > 32 else 12)
                paragraph.font.bold = row_index == 0
                paragraph.font.color.rgb = WHITE if row_index == 0 else INK


def build_presentation() -> Path:
    metrics_figure, synthetic_confusion, real_svm_confusion, real_codebert_confusion = create_figures()
    synthetic_svm = read_metrics("synthetic_svm")
    synthetic_codebert = read_metrics("synthetic_codebert")
    real_svm = read_metrics("codecontests_svm")
    real_codebert = read_metrics("codecontests_codebert_reduced")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_text(slide, "Programming Error Pattern Recognition", 0.9, 2.0, 11.6, 0.8, 38, INK, True, PP_ALIGN.CENTER)
    add_text(slide, "Pattern recognition project: prototype, real-data validation, and BIFI comparison", 1.4, 2.85, 10.6, 0.45, 18, TEAL, False, PP_ALIGN.CENTER)
    add_text(slide, "Goal: help instructors identify recurring programming mistakes without replacing human review.", 1.6, 4.0, 10.2, 0.55, 20, SLATE, False, PP_ALIGN.CENTER)

    add_slide(prs, "Project goal and PR framing", [
        "Input: a Python code submission.",
        "Output: an error-pattern label or a correct-solution label.",
        "Pattern recognition task: represent source code, train classifiers, evaluate generalization.",
        "Practical value: summarize recurring mistakes and support instructor triage.",
    ])

    add_slide(prs, "Research questions", [
        "Can source-code representations recognize recurring programming-error patterns?",
        "Does a model trained on generated labels transfer to real submissions?",
        "What does a successful research model require beyond this project setup?",
    ])

    add_slide(prs, "System built in this project", [
        "Reusable Python package with data loading, preprocessing, features, training and evaluation modules.",
        "Command-line scripts for dataset preparation, training, evaluation and prediction.",
        "Two model paths: TF-IDF + Linear SVM and CodeBERT sequence classification.",
        "Instructor-facing Streamlit dashboard for pasted code, sample examples and CSV upload.",
    ])

    add_slide(prs, "Datasets and labels", [
        "Synthetic prototype: 976 unique balanced Python snippets across eight labels.",
        "Eight-label taxonomy: correct, syntax, variable misuse, loop logic, conditional logic, function definition, data-structure misuse and algorithmic inefficiency.",
        "CodeContests validation: real Python submissions using official problem splits.",
        "Real-data labels are limited to accepted solutions and AST-confirmed syntax errors.",
    ])

    add_slide(prs, "Data integrity controls", [
        "Normalize each code sample, hash it and remove duplicates.",
        "Reject any experiment where train and test fingerprints overlap.",
        "Use a deterministic stratified split for the synthetic data.",
        "Use the official CodeContests problem split to reduce problem-level leakage.",
    ])

    add_slide(prs, "Model approaches", [
        "SVM baseline: tokenized code, unigram/bigram TF-IDF features and LinearSVC.",
        "CodeBERT: microsoft/codebert-base with a sequence-classification head.",
        "CodeBERT settings: max length 256, batch size 8, gradient accumulation, early stopping and seed 42.",
        "The full real-data CodeBERT target was resource-heavy, so a reduced CodeContests run is reported.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Measured model results", "Macro F1 is the main metric because the real test set is imbalanced")
    rows = [
        ["Model", "Dataset", "Train", "Test", "Accuracy", "Macro F1", "Weighted F1"],
        ["SVM", "Synthetic 8-class", "732", "244", f"{synthetic_svm['accuracy']:.3f}", f"{synthetic_svm['macro_f1']:.3f}", f"{synthetic_svm['weighted_f1']:.3f}"],
        ["CodeBERT", "Synthetic 8-class", "732", "244", f"{synthetic_codebert['accuracy']:.3f}", f"{synthetic_codebert['macro_f1']:.3f}", f"{synthetic_codebert['weighted_f1']:.3f}"],
        ["SVM", "CodeContests binary", "19,999", "2,059", f"{real_svm['accuracy']:.3f}", f"{real_svm['macro_f1']:.3f}", f"{real_svm['weighted_f1']:.3f}"],
        ["CodeBERT reduced", "CodeContests binary", "2,000", "2,059", f"{real_codebert['accuracy']:.3f}", f"{real_codebert['macro_f1']:.3f}", f"{real_codebert['weighted_f1']:.3f}"],
    ]
    add_table(slide, rows, 0.45, 1.45, 12.4, 2.15)
    slide.shapes.add_picture(str(metrics_figure), Inches(1.5), Inches(3.95), width=Inches(10.3))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "What the real-data results mean")
    add_bullets(slide, [
        "The CodeContests test set contains 2,000 correct solutions and only 59 syntax errors.",
        "SVM: 18/59 syntax errors found; syntax precision 0.03 and recall 0.31.",
        "Reduced CodeBERT: 51/59 syntax errors found; syntax precision 0.07 and recall 0.86.",
        "CodeBERT improves recall, but many correct programs are incorrectly flagged.",
        "The achievement is an honest validation boundary, not a claim of deployment-ready accuracy.",
    ], size=19)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Confusion matrices")
    slide.shapes.add_picture(str(synthetic_confusion), Inches(0.25), Inches(1.25), width=Inches(4.15))
    slide.shapes.add_picture(str(real_svm_confusion), Inches(4.55), Inches(1.25), width=Inches(4.15))
    slide.shapes.add_picture(str(real_codebert_confusion), Inches(8.85), Inches(1.25), width=Inches(4.15))
    add_text(slide, "Synthetic SVM is strong on template-like labels. Real-data models expose the minority-class problem clearly.", 0.8, 6.75, 11.8, 0.35, 14, SLATE, False, PP_ALIGN.CENTER)

    add_slide(prs, "Successful external study: BIFI", [
        "Break-It-Fix-It (Yasunaga and Liang, ICML 2021) is a program-repair model, not a classifier.",
        "It uses a parser/compiler-style critic to verify whether generated code is valid.",
        "Reported results: 90.5% repair accuracy on GitHub-Python AST parse errors and 71.7% on DeepFix C.",
        "It directly addresses our key limitation: synthetic errors often do not match real human errors.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Our system compared with BIFI")
    rows = [
        ["Aspect", "This project", "BIFI"],
        ["Task", "Error-pattern classification", "Syntax/compile error repair"],
        ["Training signal", "Labels from generated data and AST checks", "Parser/compiler critic plus generated pairs"],
        ["Strength", "End-to-end PR pipeline and dashboard", "High repair accuracy with real-error adaptation"],
        ["Requirement", "Labeled examples and class balance", "Large corpora, critic, more compute"],
        ["Limitation", "Low precision on real syntax errors", "Does not solve semantic bugs without stronger feedback"],
    ]
    add_table(slide, rows, 0.6, 1.45, 12.1, 4.1, header_color=TEAL)
    add_text(slide, "Future direction: use parser and unit-test feedback to move from static labels toward verified behavior.", 0.9, 6.25, 11.6, 0.5, 17, AMBER, True, PP_ALIGN.CENTER)

    add_slide(prs, "Project achievements", [
        "Built a reproducible package, scripts, tests, experiment artifacts and dashboard.",
        "Measured a strong synthetic baseline: SVM Macro F1 0.903 on eight labels.",
        "Ran real external validation on CodeContests instead of relying only on generated data.",
        "Showed that CodeBERT can improve syntax-error recall, but not precision under imbalance.",
        "Added a research comparison explaining what a successful model needs.",
    ])

    add_slide(prs, "Limitations", [
        "Synthetic eight-class labels demonstrate feasibility but not classroom generalization.",
        "CodeContests is competitive-programming data, not a classroom dataset.",
        "Real semantic labels are unavailable; parseable incorrect submissions are excluded.",
        "The current task is single-label, while real code can contain multiple errors.",
        "A reliable production model needs expert annotation or test-based behavioral labels.",
    ])

    add_slide(prs, "Conclusion", [
        "The project achieved a complete PR workflow for programming-error pattern recognition.",
        "The synthetic task shows the pipeline can learn the proposed taxonomy.",
        "The real-data task shows why accuracy alone is misleading under severe imbalance.",
        "BIFI shows the next step: combine large data with parser/compiler or test feedback.",
    ])

    add_slide(prs, "References", [
        "Yasunaga and Liang, Break-It-Fix-It: Unsupervised Learning for Program Repair, ICML 2021.",
        "Yasunaga and Liang, Graph-based, Self-Supervised Program Repair from Diagnostic Feedback, ICML 2020.",
        "DeepMind CodeContests dataset for real competitive-programming submissions.",
        "Microsoft CodeBERT for transformer-based source-code representation.",
    ])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_presentation())
